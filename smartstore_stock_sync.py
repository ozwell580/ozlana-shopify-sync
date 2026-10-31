"""
스마트스토어 재고 업로드 파일 생성 (DK UGG + 에버어그 + 오즈웨어 + 오즈라나)
------------------------------------------------------------------------
입력
  forms/ 폴더 안의 모든 .xlsx
    - 스마트스토어 '수정양식 다운로드'로 받은 파일들 (DK, 에버어그, 키즈 등 여러 개 가능)
    - 브랜드 칸(DK어그 / 에버어그 / 오즈웨어 / 오즈라나)을 보고 자동으로 나눠서 처리합니다.
      브랜드 칸이 비었거나 다르게 적혀 있으면 상품명(오즈웨어/오즈라나)이나 코드(DK...)로 판단합니다.
  공급처 재고 (실행할 때 자동 호출)
    - DK UGG 재고 API
    - 에버어그: GitHub 시크릿 EVERUGG_USER_ID / EVERUGG_PASSWORD
    - 오즈라나: GitHub 시크릿 OZLANA_TOKEN (오즈라나 /stocks API)
    - 오즈웨어: GitHub 시크릿 OZWEAR_API_KEY / OZWEAR_API_SECRET + SHOPIFY_STORE / SHOPIFY_CLIENT_ID / SHOPIFY_CLIENT_SECRET
      (Shopify 'OZWEAR UGG' 상품의 SKU·색상·사이즈를 읽고, 그 SKU로 오즈웨어 API 재고를 조회)
출력
  smartstore_stock_upload.xlsx : 스마트스토어 '수정양식 업로드'에 올리는 파일
  smartstore_stock_report.txt  : 처리 결과 보고서

규칙
  - DK: 사이즈 옵션을 DK 사이즈표(EU)로 맞춤. 남는 옛 옵션은 숨김(N), 부족한 옵션은 '추가'
  - 에버어그: 사이즈 표기는 그대로 두고 에버어그 재고로 갱신
  - 공급처에 없는 색상/사이즈/상품 -> 재고 0 (품절)  [DK·에버어그]
  - 오즈웨어·오즈라나는 공급처에서 못 찾은 상품/색상/사이즈의 재고를 그대로 둠 (재고 유지, 보고서에 표시)
  - 오즈라나 사이즈는 AU 사이즈 -> EU(AU+30, 예: AU5 -> EU35)로 바꿔 맞춤. 바꾸려면 OZ_EU_OFFSET
  - 단독형(컬러만) -> 조합형으로 바꿔서 색상별 재고 입력
  - 간편 입력 옵션 상품(옵션형태 공란)은 엑셀 수정이 안 돼서 업로드 파일에서 제외
  - 판매자 상품코드가 빈 상품은 상품명 속 코드로 채움
  - 같은 상품번호가 여러 파일에 있으면 한 번만 처리
  - 유료배송(1개마다 부과): 판매가에서 배송비 몫을 빼고 배송비로 받음 (롱부츠·가방·아우터 17,000 / 미니·숏·앵클부츠·플랫폼·스니커즈 14,000 / 키링·참·펜던트 9,900 / 그 외 11,000원)
    이미 바뀐 상품은 다시 빼지 않음. 끄려면 PAID_SHIPPING=0
"""

import glob
import json
import os
import re
import sys
import time
import collections
import requests
import openpyxl

FORMS_DIR = os.environ.get("FORMS_DIR", "forms")
OUT_FILE = "smartstore_stock_upload.xlsx"
UPLOAD_LIMIT = 450  # 스마트스토어 수정양식 업로드 최대 500개
REPORT_FILE = "smartstore_stock_report.txt"

SAFETY_STOCK = int(os.environ.get("SAFETY_STOCK", "0"))  # 공급처 재고에서 빼고 올릴 수량
MAX_STOCK = 999
USE_CACHE = os.environ.get("USE_CACHE") == "1"
# 어린이인증 상품에 넣을 상품정보제공고시 템플릿 코드 (비워두면 넣지 않음)
# 템플릿 방식은 통과되지 않아서 기본값은 비워둠. 키즈 상품은 스마트스토어에서 모델명을 직접 입력.
KIDS_NOTICE_TEMPLATE = os.environ.get("KIDS_NOTICE_TEMPLATE", "")
KIDS_RE = re.compile(r"키즈|아동|유아|베이비|토들러|인형")

DK_URLS = [
    ("GET", "https://api.dkugg.dreamdaytech.com/api/product_stock"),
    ("POST", "https://api.dkugg.dreamdaytech.com/api/product_stock"),
    ("GET", "http://api.dkugg.dreamdaytech.com/api/product_stock"),
]

EVERUGG_BASE_URL = os.environ.get("EVERUGG_BASE_URL") or "http://api.everugg.net.au:9990"
EVERUGG_USER_ID = os.environ.get("EVERUGG_USER_ID")
EVERUGG_PASSWORD = os.environ.get("EVERUGG_PASSWORD")
EVERUGG_WAREHOUSES = ["/Api/Stock/AuStock", "/Api/Stock/SydStock", "/Api/Stock/SydrhStock"]

OZLANA_BASE_URL = os.environ.get("OZLANA_BASE_URL") or "http://www.ozlanacms.com.au:30008"
OZLANA_TOKEN = os.environ.get("OZLANA_TOKEN")
OZWEAR_BASE_URL = "https://api.ozwearugg.net/rest/s1/openapi"
OZWEAR_API_KEY = os.environ.get("OZWEAR_API_KEY")
OZWEAR_API_SECRET = os.environ.get("OZWEAR_API_SECRET")
OZWEAR_VENDOR = "OZWEAR UGG"
SHOPIFY_STORE = os.environ.get("SHOPIFY_STORE")
SHOPIFY_CLIENT_ID = os.environ.get("SHOPIFY_CLIENT_ID")
SHOPIFY_CLIENT_SECRET = os.environ.get("SHOPIFY_CLIENT_SECRET")
OZ_EU_OFFSET = int(os.environ.get("OZ_EU_OFFSET", "30"))  # AU5 + 30 = EU35

BRAND_DK = "DK어그"
BRAND_EU = "에버어그"
BRAND_OW = "오즈웨어"
BRAND_OL = "오즈라나"
OZ_BRANDS = (BRAND_OW, BRAND_OL)
SUPPLIER_NOTES = []  # 공급처 조회 문제 (보고서 맨 위에 표시)

# 한글 색상 -> 영문 (띄어쓰기 제거한 한글 기준, DK/에버어그 공용)
COLOR_MAP = {
    "그레이": "Grey", "그린": "Green", "누드": "Nude", "네이비블루": "Navy Blue",
    "다크그레이": "Dark Grey", "다크브라운": "Dark Brown", "던핑크": "Dawn Pink",
    "데님블루": "Denim Blue", "라벤더": "Lavender", "라이트그레이": "Light Grey",
    "라이트블루": "Light Blue", "스카이블루": "Skyblue", "라이트브라운": "Light Brown",
    "라이트카멜": "Light Camel", "라이트핑크": "Light Pink", "라일락": "Lilac",
    "레드": "Red", "레몬옐로우": "Lemon Yellow", "레오파드": "Leopard", "레인보우": "Rainbow",
    "로즈브라운": "Rosy Brown", "로지브라운": "Rosy Brown", "로즈핑크": "Rose Pink",
    "로지던": "Rosy Dawn", "로지핑크": "Rosy Pink", "롤랜드그레이": "Roland Grey",
    "머스타드씨드": "Mustard Seed", "메이플레드": "Maple Red", "모스그린": "Moss Green",
    "미스트핑크": "Mist Pink", "밀키화이트": "Milky White", "밸리브라운": "Vally Brown",
    "밸리브라운그린": "Vally Brown/Green", "밸리브라운핑크": "Vally Brown/Pink",
    "베이지": "Beige", "베이지화이트": "Beige White", "브라운": "Brown", "브론즈": "Bronze",
    "브릭그레이": "Brick Gray", "블랙": "Black", "블랙핑크": "Black Pink", "블루": "Blue",
    "블루블랙": "Blue Black", "블러쉬핑크": "Blush Pink", "블러시핑크": "Blush Pink",
    "상그리아": "Sangria", "샌드": "Sand", "샌드스톤": "Sandstone", "샌드라임": "Sand/Lime",
    "샌드아쿠아": "Sand/Aqua", "샴브레이": "Chambray", "스플렌디드": "Splendid", "시슬": "Thistle",
    "실버": "Silver", "아이보리": "Ivory", "아이보리핑크": "Ivory/Pink", "아쿠아블루": "Aqua Blue",
    "애쉬브라운": "Ash Brown", "애프리콧": "Apricot", "애플그린": "Apple Green",
    "앤틸로프브라운": "Antelope Brown", "앤틸롭브라운": "Antelope Brown", "엔틸롭브라운": "Antelope Brown",
    "어도비그레이": "Adobe Grey", "엄버": "Umber", "에어리블루": "Airy Blue", "에크루": "Ecru",
    "옐로우": "Yellow", "오렌지": "Orange", "오크": "Oak", "오트밀": "Oatmeal",
    "오프화이트": "Off White", "올리브": "Olive", "올리브브라운": "Olive Brown",
    "옵시디언그레이": "Obsidian Gray", "와인": "Wine", "우드랜드그린": "Woodland Green",
    "월넛브라운": "Walnut Brown", "웜그레이": "Warm Grey", "유칼립투스": "Eucalyptus",
    "잉크네이비": "Ink Navy", "차콜그레이": "Charcoal Grey", "체스트넛": "Chestnut",
    "체스넛": "Chestnut", "체스트넛레드": "Chestnut/Red", "체스트넛베이지": "Chestnut-Beige",
    "체스트넛브라운": "Chestnut/Brown", "체스트넛크림": "Chestnut Cream",
    "체스트넛모스그린": "Chestnut Moss green", "초콜릿": "Chocolate", "초콜렛": "Chocolate",
    "카놀라": "Canola", "카라멜": "Caramel", "카라멜브라운": "Caramel Brown", "카리부": "Caribou",
    "칼리보": "Caribou", "카멜": "Camel", "카키": "Khaki", "커스터드": "Custard",
    "커피": "Coffee", "컬러블록블루": "Colorblock Blue", "코코아": "Cocoa", "쿠퍼": "Copper",
    "크림": "Cream", "크림그레이": "Cream/Grey", "크림밀크": "Cream/Milk",
    "크림카리부": "Cream/Caribou", "크림체스트넛": "Cream/Chestnut", "크림화이트": "Cream/White",
    "클레이브라운": "Clay Brown", "토프": "Taupe", "페블그레이": "Pebble Grey",
    "포그그레이": "Fog Grey", "프라신그린": "Prasine", "피치": "Peach", "피치핑크": "Peach Pink",
    "피칸": "Pecan", "핑크": "Pink", "헤이즐넛": "Hazelnut", "화이트": "White",
    "고트그레이": "Goat Grey", "나파블랙": "nappa black", "나파초콜릿": "nappa chocolate",
    "내추럴아이보리": "natural lvory",
    # 체크 패턴
    "그레이체크": "Grey Check", "그레이핑크체크": "Grey Pink Check", "네이비체크": "Navy Check",
    "라이트블루체크": "Light Blue Check", "라이트핑크체크": "Light Pink Check",
    "레드체크": "Red Check", "베이지체크": "Beige Check", "블랙체크": "Black Check",
    "인디고카멜체크": "Indigo Camel Check", "카멜체크": "camel check", "카키체크": "Khaki Check",
    # 동물 키링
    "곰발바닥": "Bear Paw", "북극곰": "Polar Bear", "판다": "Panda", "코알라": "Koala",
    "캥거루": "Kangaroo", "회색양": "Grey Sheep", "흰색양": "White Sheep",
    # 오즈웨어·오즈라나
    "다크초콜렛": "Dark Chocolate", "다크초콜릿": "Dark Chocolate", "아이보리블랙": "Ivory/Black",
    "탄": "Tan", "히코리": "Hickory", "머쉬룸": "Mushroom", "머시룸": "Mushroom",
}

KIDS_EU_TO_AU = {"21-22": "4-5", "23-25": "6-7", "26-28": "8-10", "29-32": "11-12", "33-34": "13-2"}
BABY_MONTHS = {"S": "0-6 M", "M": "6-12 M", "L": "12-18 M", "XL": "18-24 M"}

CODE_RE = re.compile(r"\b(D[A-Z]\d+[A-Z]?|AS[A-Z]?\d+[A-Z]*|OZB\d+[A-Z]*|OZ\d+[A-Z]*|OB\d+[A-Z]*|CS\d+[A-Z]*|\d{5})\b")

# ================================================================ 유료배송 (1개마다 배송비 부과)
# 판매가에서 배송비만큼 빼고 그 금액을 배송비로 받음 -> 고객 결제 총액은 그대로
#   17,000원 : 롱·미들 부츠, 가방, 아우터(자켓·코트·무스탕·베스트·패딩)
#   14,000원 : 미니·숏·앵클 부츠(남성·키즈 포함), 플랫폼(슬리퍼·샌들·로퍼 등), 스니커즈, 운동화
#   11,000원 : 그 외 (일반 슬리퍼·샌들·모자·장갑·머플러·부티·의류 등)
#    9,900원 : 키링, 가방 참, 펜던트
PAID_SHIPPING = os.environ.get("PAID_SHIPPING", "1") == "1"
SHIP_FEE_BIG = 17000
SHIP_FEE_MID = 14000
SHIP_FEE_LOW = 11000
SHIP_FEE_CHARM = 9900
SHIP_CHARM = re.compile(r"키링|백참|가방참|가방 참|펜던트")
SHIP_CHARM_WORD = re.compile(r"참\b")
SHIP_SHOE_WORD = re.compile(r"로퍼|뮬|슬리퍼|샌들|부츠|스니커즈|클로그|슈즈|슬립온|메리제인")
SHIP_LOW_STRONG = re.compile(r"관리|스프레이|양말|삭스|부티|더스트백|신발주머니|부츠커버|레그워머")
SHIP_BOOT = re.compile(r"부츠|부처|숏어그|롱어그|아웃도어어그|워커")
SHIP_LONG_BOOT = re.compile(r"롱|미들|미디|톨\s|톨부츠|니하이|하이\s?부츠")
SHIP_OUTER = re.compile(r"자켓|재킷|코트|점퍼|패딩|베스트|조끼")
SHIP_MID = re.compile(r"플랫폼|스니커즈|운동화")
SHIP_BAG = re.compile(r"가방|백팩|(숄더|크로스|토트|버킷|호보|미니|클러치|체인|새들|쇼퍼|더플|핸들|바디)\s?백|클러치|토트")
SHIP_LOW_ACC = re.compile(r"모자|햇|장갑|미튼|머플러|목도리|귀마개|비니|캡\b|볼캡")


def ship_fee_for(name):
    n = str(name or "")
    if SHIP_CHARM.search(n) or (SHIP_CHARM_WORD.search(n) and not SHIP_SHOE_WORD.search(n)):
        return SHIP_FEE_CHARM
    if SHIP_LOW_STRONG.search(n):
        return SHIP_FEE_LOW
    if SHIP_BOOT.search(n):
        if "미니" in n or not SHIP_LONG_BOOT.search(n):
            return SHIP_FEE_MID
        return SHIP_FEE_BIG
    if SHIP_OUTER.search(n):
        return SHIP_FEE_BIG
    if SHIP_MID.search(n):
        return SHIP_FEE_MID
    if SHIP_LOW_ACC.search(n):
        return SHIP_FEE_LOW
    if SHIP_BAG.search(n):
        return SHIP_FEE_BIG
    return SHIP_FEE_LOW


def apply_paid_shipping(r, col, R):
    """무료/유료 배송 상품을 '수량별(1개마다)' 유료배송으로 바꾸고 판매가에서 배송비만큼 뺌.
    기존 배송비는 판매가에 다시 더한 뒤 새 배송비를 빼므로 고객 총액은 그대로.
    이미 기준 금액으로 바뀐 상품은 그대로 둠."""
    typ = str(r[col["배송비유형"]] or "").strip()
    old_fee = to_int(r[col["기본배송비"]])
    fee = ship_fee_for(r[col["상품명"]])
    if typ == "수량별" and old_fee == fee and str(r[col["수량별부과-수량"]] or "").strip() in ("1", "1.0"):
        return
    if typ not in ("무료", "유료", "수량별"):
        R["배송비 유형 확인 필요(변경 안 함)"].append(f"{r[col['상품번호']]}  {typ}  {r[col['상품명']]}")
        return
    if typ == "무료":
        old_fee = 0
    price = to_int(r[col["판매가"]])
    new_price = price + old_fee - fee
    if new_price < 1000:
        R["배송비 유형 확인 필요(변경 안 함)"].append(f"{r[col['상품번호']]}  판매가 {price}원이 너무 낮음  {r[col['상품명']]}")
        return
    r[col["판매가"]] = new_price
    r[col["배송비유형"]] = "수량별"
    r[col["기본배송비"]] = fee
    r[col["배송비 결제방식"]] = "선결제"
    r[col["조건부무료-\n상품판매가 합계"]] = None
    r[col["수량별부과-수량"]] = 1
    R["유료배송 전환"].append(f"{r[col['판매자 상품코드']]}  판매가 {price:,} -> {new_price:,}  배송비 {fee:,}"
                         + (f" (기존 배송비 {old_fee:,})" if old_fee else "") + f"  {r[col['상품명']]}")


# ================================================================ 공급처 데이터
def to_int(v):
    try:
        return max(int(float(v)), 0)
    except Exception:
        return {"sufficient": 50, "shortage": 3, "out": 0}.get(str(v).strip().lower(), 0)


def fetch_dk():
    if USE_CACHE and os.path.exists("dk_stock.json"):
        with open("dk_stock.json", encoding="utf-8") as f:
            return json.load(f)["products"]
    for attempt in range(3):
        for method, url in DK_URLS:
            try:
                if method == "GET":
                    r = requests.get(url, params={"num": 1}, timeout=120)
                else:
                    r = requests.post(url, params={"num": 1}, json={},
                                      headers={"content-type": "application/json"}, timeout=120)
                if r.status_code == 200:
                    data = r.json()
                    if isinstance(data, dict) and isinstance(data.get("products"), list):
                        print(f"DK API 성공 [{method}] {url}")
                        with open("dk_stock.json", "w", encoding="utf-8") as f:
                            json.dump(data, f, ensure_ascii=False)
                        return data["products"]
                print(f"DK API [{method}] {url} -> {r.status_code}")
            except Exception as e:
                print(f"DK API [{method}] {url} 실패: {e}")
        time.sleep(10 * (attempt + 1))
    sys.exit("DK API 호출 실패")


def everugg_token():
    url = f"{EVERUGG_BASE_URL}/Api/Token/getToken"
    for attempt in range(3):
        try:
            r = requests.get(url, params={"user": EVERUGG_USER_ID, "password": EVERUGG_PASSWORD}, timeout=30)
            if r.status_code == 200:
                res = r.json().get("result", {})
                tok = res.get("token") if isinstance(res, dict) else None
                if tok:
                    return tok
            print(f"에버어그 토큰 실패 ({attempt + 1}차): {r.status_code}")
        except Exception as e:
            print(f"에버어그 토큰 예외 ({attempt + 1}차): {e}")
        time.sleep(10 * (attempt + 1))
    sys.exit("에버어그 토큰 발급 실패")


def everugg_list(path, token):
    url = f"{EVERUGG_BASE_URL.rstrip('/')}{path}"
    for attempt in range(3):
        try:
            r = requests.get(url, params={"token": token}, headers={"Accept": "application/json"}, timeout=90)
            r.raise_for_status()
            body = r.json()
            if isinstance(body, list):
                return body
            if isinstance(body, dict):
                for k in ("result", "data", "items", "list"):
                    v = body.get(k)
                    if isinstance(v, list):
                        return v
                    if isinstance(v, dict):
                        for k2 in ("data", "items", "list", "result"):
                            if isinstance(v.get(k2), list):
                                return v[k2]
            return []
        except Exception as e:
            print(f"{path} 실패 ({attempt + 1}차): {e}")
            time.sleep(10 * (attempt + 1))
    sys.exit(f"에버어그 {path} 조회 실패")


def fetch_everugg():
    if USE_CACHE and os.path.exists("everugg_stock.json"):
        with open("everugg_stock.json", encoding="utf-8") as f:
            return json.load(f)
    if not EVERUGG_USER_ID or not EVERUGG_PASSWORD:
        sys.exit("EVERUGG_USER_ID / EVERUGG_PASSWORD 시크릿이 없습니다.")
    token = everugg_token()
    combined = {}
    for path in EVERUGG_WAREHOUSES:
        rows = everugg_list(path, token)
        print(f"에버어그 {path}: {len(rows)}건")
        for it in rows:
            if not isinstance(it, dict):
                continue
            bc = str(it.get("Barcode", "")).strip()
            if not bc:
                continue
            q = int(it.get("AvaiStockQty", 0) or 0)
            if bc in combined:
                combined[bc]["AvaiStockQty"] += q
            else:
                combined[bc] = dict(it)
                combined[bc]["AvaiStockQty"] = q
    items = list(combined.values())
    with open("everugg_stock.json", "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False)
    return items


def oz_size(size):
    """오즈 사이즈 -> 비교용 표기. AU 숫자(3~13)는 EU로 바꿈 (AU5 -> EU35)."""
    s = str(size or "").strip()
    u = s.upper().replace(" ", "")
    if "#" in u:  # 오즈라나 '5#...' 형식
        parts = [x for x in u.split("#") if x]
        eu = [x for x in parts if re.fullmatch(r"(EU)?(3\d|4\d)", x)]
        if eu:
            return "EU" + eu[0].replace("EU", "")
        u = parts[0] if parts else ""
    m = re.fullmatch(r"(AU|W|US)?(\d{1,2})", u)
    if m and 3 <= int(m.group(2)) <= 13 and m.group(1) != "US":
        return f"EU{int(m.group(2)) + OZ_EU_OFFSET}"
    m = re.fullmatch(r"EU(\d+)", u)
    if m:
        return u
    return s


def fetch_ozlana():
    """{상품코드: [ {color, size, stock} ]}  (오즈라나 /stocks API)"""
    if not OZLANA_TOKEN:
        SUPPLIER_NOTES.append("오즈라나: OZLANA_TOKEN 시크릿이 없어 재고를 가져오지 못함 (오즈라나 재고 유지)")
        return {}
    for attempt in range(3):
        try:
            r = requests.get(f"{OZLANA_BASE_URL}/stocks", headers={"X-Token": OZLANA_TOKEN}, timeout=120)
            if r.status_code == 200:
                raw = r.json().get("data", [])
                prods = json.loads(raw) if isinstance(raw, str) else raw
                out = collections.defaultdict(list)
                for p in prods or []:
                    if not isinstance(p, dict):
                        continue
                    mark = str(p.get("prodMark", "")).strip().upper()
                    if not mark:
                        continue
                    color = str(p.get("colorName", "")).strip()
                    for st in p.get("stocks") or []:
                        if not isinstance(st, dict):
                            continue
                        item = {"color": color, "size": oz_size(st.get("sizeRemark")),
                                "raw_size": str(st.get("sizeRemark", "")), "stock": to_int(st.get("stockNum"))}
                        keys = {mark}
                        if mark.startswith("OZ3"):
                            keys.add("OZ" + mark[3:])
                        for k in keys:
                            out[k].append(item)
                print(f"오즈라나 재고: {len(out)}개 상품코드")
                return out
            print(f"오즈라나 API -> {r.status_code}")
        except Exception as e:
            print(f"오즈라나 API 실패 ({attempt + 1}차): {e}")
        time.sleep(10 * (attempt + 1))
    SUPPLIER_NOTES.append("오즈라나: API 호출 실패 (오즈라나 재고 유지)")
    return {}


def shopify_token():
    if not SHOPIFY_STORE or not SHOPIFY_CLIENT_ID or not SHOPIFY_CLIENT_SECRET:
        return None, None
    domain = SHOPIFY_STORE.replace("https://", "").strip("/")
    r = requests.post(f"https://{domain}/admin/oauth/access_token", json={
        "client_id": SHOPIFY_CLIENT_ID, "client_secret": SHOPIFY_CLIENT_SECRET,
        "grant_type": "client_credentials"}, timeout=30)
    if r.status_code != 200:
        print(f"Shopify 토큰 실패: {r.status_code}")
        return None, domain
    return r.json().get("access_token"), domain


def fetch_ozwear():
    """{상품코드: [ {color, size, stock} ]}
    Shopify 'OZWEAR UGG' 상품에서 SKU·색상·사이즈를 읽고, 오즈웨어 API로 SKU별 재고 조회"""
    if not OZWEAR_API_KEY or not OZWEAR_API_SECRET:
        SUPPLIER_NOTES.append("오즈웨어: OZWEAR_API_KEY / OZWEAR_API_SECRET 시크릿이 없어 재고를 가져오지 못함 (오즈웨어 재고 유지)")
        return {}
    try:
        token, domain = shopify_token()
        if not token:
            SUPPLIER_NOTES.append("오즈웨어: Shopify 접속 실패 (SHOPIFY_* 시크릿 확인, 오즈웨어 재고 유지)")
            return {}
        headers = {"X-Shopify-Access-Token": token}
        url = (f"https://{domain}/admin/api/2024-01/products.json"
               f"?limit=250&vendor={requests.utils.quote(OZWEAR_VENDOR)}")
        variants = []  # (codes, sku, color, size)
        while url:
            r = requests.get(url, headers=headers, timeout=60)
            if r.status_code != 200:
                raise RuntimeError(f"Shopify 상품 조회 {r.status_code}")
            for p in r.json().get("products", []):
                names = [str(o.get("name", "")).lower() for o in p.get("options") or []]
                c_i = next((i for i, n in enumerate(names) if "colo" in n or "색" in n), None)
                s_i = next((i for i, n in enumerate(names) if "size" in n or "사이즈" in n), None)
                title_codes = {m.upper() for m in CODE_RE.findall(str(p.get("title", "")).upper())}
                for v in p.get("variants") or []:
                    sku = str(v.get("sku") or "").strip()
                    if not sku:
                        continue
                    opts = [v.get("option1"), v.get("option2"), v.get("option3")]
                    color = (opts[c_i] if c_i is not None and c_i < 3 else None) or ""
                    size = opts[s_i] if s_i is not None and s_i < 3 else None
                    codes = set(title_codes)
                    head = re.split(r"[-_ ]", sku.upper())[0]
                    if head:
                        codes.add(head)
                    variants.append((codes, sku, color, size))
            url = None
            for link in (r.headers.get("Link") or "").split(","):
                if 'rel="next"' in link:
                    url = link.split(";")[0].strip("<> ")
        print(f"Shopify 오즈웨어 SKU: {len(variants)}개")
        if not variants:
            SUPPLIER_NOTES.append("오즈웨어: Shopify에 'OZWEAR UGG' 상품이 없음 (오즈웨어 재고 유지)")
            return {}

        r = requests.post(f"{OZWEAR_BASE_URL}/token", json={"key": OZWEAR_API_KEY, "secret": OZWEAR_API_SECRET}, timeout=30)
        r.raise_for_status()
        tok = r.json()["data"]["token"]
        skus = sorted({v[1] for v in variants})
        stock = {}
        for i in range(0, len(skus), 50):
            r = requests.post(f"{OZWEAR_BASE_URL}/products/stock",
                              headers={"api_key": tok, "Content-Type": "application/json"},
                              json={"skus": skus[i:i + 50], "productIds": []}, timeout=60)
            r.raise_for_status()
            for it in (r.json().get("data") or {}).get("list") or []:
                k = str(it.get("sku", "")).strip()
                if k:
                    stock[k] = to_int(it.get("stock"))
            time.sleep(0.3)
        print(f"오즈웨어 API 재고: {len(stock)}개 SKU")
        out = collections.defaultdict(list)
        for codes, sku, color, size in variants:
            if sku not in stock:
                continue
            item = {"color": color, "size": oz_size(size) if size is not None else "FREE",
                    "raw_size": str(size), "stock": stock[sku]}
            for c in codes:
                out[c].append(item)
        return out
    except Exception as e:
        print(f"오즈웨어 재고 조회 실패: {e}")
        SUPPLIER_NOTES.append(f"오즈웨어: 재고 조회 실패 ({e}) (오즈웨어 재고 유지)")
        return {}


def brand_of(r, col):
    """브랜드 칸이 비었거나 다르게 적힌 상품도 상품명/코드로 판단"""
    b = str(r[col["브랜드"]] or "").strip()
    name = str(r[col["상품명"]] or "")
    code = str(r[col["판매자 상품코드"]] or "").upper()
    if b in (BRAND_DK, BRAND_EU, BRAND_OW, BRAND_OL):
        return b
    if "오즈라나" in b or "오즈라나" in name:
        return BRAND_OL
    if "오즈웨어" in b or "오즈웨어" in name:
        return BRAND_OW
    if "에버어그" in b or "에버어그" in name:
        return BRAND_EU
    m = CODE_RE.search(name.upper())
    c = code or (m.group(1) if m else "")
    if b in ("디케이", "DK", "DK UGG", "어그", "") and re.fullmatch(r"D[A-Z]\d+[A-Z]?", c):
        return BRAND_DK
    return b or None


def build_supplier(brand_needed):
    """{브랜드: {상품코드: [ {color, size, stock}, ... ]}}"""
    sup = {}
    if BRAND_DK in brand_needed:
        dk = collections.defaultdict(list)
        for p in fetch_dk():
            for s in p.get("skus") or []:
                dk[str(p.get("id")).strip().upper()].append(
                    {"color": s.get("color"), "size": str(s.get("size")), "stock": to_int(s.get("stock"))})
        sup[BRAND_DK] = dk
    if BRAND_EU in brand_needed:
        eu = collections.defaultdict(list)
        for it in fetch_everugg():
            eu[str(it.get("ProductCode", "")).strip().upper()].append(
                {"color": it.get("ColorName"), "size": str(it.get("Size")), "stock": to_int(it.get("AvaiStockQty"))})
        sup[BRAND_EU] = eu
    if BRAND_OL in brand_needed:
        sup[BRAND_OL] = fetch_ozlana()
    if BRAND_OW in brand_needed:
        sup[BRAND_OW] = fetch_ozwear()
    return sup


def final_qty(q):
    return min(max(q - SAFETY_STOCK, 0), MAX_STOCK)


# ================================================================ 색상
def norm_ko(s):
    return re.sub(r"[\s/]+", "", str(s or ""))


def norm_en(s):
    return re.sub(r"[\s\-/()&]+", "", str(s or "")).lower()


def find_color(skus, ko, allow_guess=True, allow_single=True):
    """한글 색상 -> 공급처 색상. (색상, 추정여부)
    allow_single: 공급처 색상이 하나뿐이면 그 색으로 봄 (오즈는 영문이 다르면 쓰지 않음)"""
    colors = sorted({str(s["color"]) for s in skus})
    en = COLOR_MAP.get(norm_ko(ko))
    if en:
        for c in colors:
            if norm_en(c) == norm_en(en):
                return c, False
        # 예: 그레이 -> Pebble Grey 처럼 포함되는 색상이 하나뿐이면 추정 매칭
        part = [c for c in colors if norm_en(en) in norm_en(c)]
        if allow_guess and len(part) == 1:
            return part[0], True
    if len(colors) == 1 and (allow_single or not en):
        return colors[0], False
    return None, False


# ================================================================ 사이즈
def api_size_keys(size):
    s = str(size or "").strip().upper().replace("\\", "/").replace(" ", "")
    keys = {s}
    if s in ("", "NONE", "ONESIZE", "FREESIZE", "N/A"):
        keys.add("FREE")
    m = re.search(r"\(([\d/]+)\)", s)  # DK: 3M/5L(36/37)
    if m:
        for n in m.group(1).split("/"):
            keys.add("EU" + n)
        m2 = re.search(r"(\d+)L", s)
        if m2:
            keys.add("AU" + m2.group(1))
        head = re.sub(r"\(.*", "", s)
        if head and not re.search(r"\d", head):
            keys.add(head)
    if re.fullmatch(r"\d+(/\d+)+", s):  # 35/36
        for n in s.split("/"):
            keys.add("EU" + n)
    if re.fullmatch(r"\d+(\.\d)?", s):  # 35, 7.5
        keys.add("EU" + s)
    keys.add(s.replace("-", "/"))
    return keys


def store_size_keys(label):
    s = str(label or "").strip().upper().replace(" ", "")
    keys = set()
    head = re.sub(r"\(.*", "", s)
    inner = re.search(r"\((.*)\)", s)
    keys.add(head)
    keys.add(head.replace("-", "/"))
    if re.fullmatch(r"AU\d+", head):
        keys.add(head)
    m = re.fullmatch(r"EU(\d+)", head)
    if m:
        keys.add("EU" + m.group(1))
        keys.add(m.group(1))
    m = re.fullmatch(r"EU(\d+)-(\d+)", head)
    if m:
        rng = f"{m.group(1)}-{m.group(2)}"
        keys.update({rng, rng.replace("-", "/")})
        if rng in KIDS_EU_TO_AU:
            keys.add(KIDS_EU_TO_AU[rng])
    if re.fullmatch(r"2\d\d", head):  # 225 (mm) -> EU35
        eu = 35 + (int(head) - 225) // 5
        keys.update({str(eu), f"EU{eu}"})
    if re.fullmatch(r"\d+(\.\d)?", head):
        keys.add("EU" + head)
    if inner and "개월" in inner.group(1) and head in BABY_MONTHS:
        keys.add(BABY_MONTHS[head].replace(" ", ""))
    if head in ("FREE", "FREESIZE", "ONESIZE"):
        keys.add("FREE")
    return keys


def is_size_label(v):
    s = str(v or "").strip().upper().replace(" ", "")
    head = re.sub(r"\(.*", "", s)
    return bool(re.fullmatch(
        r"(AU\d+|EU\d+(-\d+)?|XXL|XL|L|M|S|XS|S/M|M/L|L/XL|S-M|L-XL|FREE|FREESIZE|ONESIZE|L\d+|M\d+|2\d\d|\d+(\.\d)?|\d+-\d+)",
        head))


def size_stock(skus, color, label):
    keys = store_size_keys(label)
    hits = [s for s in skus if (color is None or s["color"] == color) and keys & api_size_keys(s["size"])]
    if not hits:
        pool = [s for s in skus if color is None or s["color"] == color]
        if len(pool) == 1 and "FREE" in api_size_keys(pool[0]["size"]):
            hits = pool
    return (sum(s["stock"] for s in hits), True) if hits else (0, False)


# ================================================================ DK 사이즈 재구성용
def eu_label(api_size):
    s = str(api_size or "").strip()
    m = re.search(r"\(([\d/]+)\)", s)
    if not m:
        return s
    head = re.sub(r"\(.*", "", s).strip()
    if head and not re.search(r"\d", head):
        return f"{head}(EU{m.group(1)})"
    return f"EU{m.group(1)}"


def eu_sort_key(api_size):
    m = re.search(r"\((\d+)", str(api_size or ""))
    return int(m.group(1)) if m else 999


def size_sort_key(size):
    s = str(size or "").upper()
    m = re.search(r"\((\d+)", s)
    if m:
        return (0, int(m.group(1)), s)
    m = re.fullmatch(r"([LM])(\d+)", s)
    if m:
        return (1 if m.group(1) == "L" else 2, int(m.group(2)), s)
    m = re.match(r"(\d+(\.\d)?)", s)
    if m:
        return (0, float(m.group(1)), s)
    order = {"XS": 1, "S": 2, "S-M": 3, "S/M": 3, "M": 4, "M/L": 5, "L": 6, "L-XL": 7, "XL": 8, "XXL": 9, "2XL": 9}
    return (3, order.get(s, 99), s)


def everugg_label(api_size):
    """에버어그 사이즈 -> 스마트스토어 표기"""
    s = str(api_size or "").strip()
    u = s.upper().replace(" ", "")
    if re.fullmatch(r"\d+", u):
        return f"EU{u}" if int(u) >= 16 else u
    if re.fullmatch(r"\d+\\\d+", u):
        return "EU" + u.replace("\\", "/")
    m = re.fullmatch(r"L(\d+)", u)
    if m:
        return f"L{m.group(1)}(EU{int(m.group(1)) + 31})"
    m = re.fullmatch(r"M(\d+)", u)
    if m:
        return f"M{m.group(1)}(EU{int(m.group(1)) + 34})"
    rev = {v: k for k, v in KIDS_EU_TO_AU.items()}
    if u in rev:
        return f"EU{rev[u]}"
    if u in ("ONESIZE", "NONE", ""):
        return "FREE"
    return s


# ================================================================ 메인
def main():
    files = sorted(glob.glob(os.path.join(FORMS_DIR, "*.xlsx")))
    if not files:
        sys.exit(f"{FORMS_DIR} 폴더에 수정양식 파일이 없습니다.")
    print("수정양식 파일:", files)

    rows, header = [], None
    seen = set()
    dup_names = collections.Counter()
    for f in files:
        ws = openpyxl.load_workbook(f).active
        h = [c.value for c in ws[2]]
        if header is None:
            header = h
        elif h != header:
            print(f"⚠️ {f} 양식 칸 구성이 달라서 건너뜁니다.")
            continue
        for r in ws.iter_rows(min_row=6, values_only=True):
            if not r[0] or str(r[0]) in seen:
                continue
            seen.add(str(r[0]))
            rows.append(list(r))
    col = {k: i for i, k in enumerate(header)}
    need = ["상품번호", "판매자 상품코드", "상품명", "브랜드", "재고수량", "옵션형태", "옵션명",
            "옵션번호", "옵션값", "옵션가", "옵션 재고수량", "옵션 사용여부", "상품정보제공고시 템플릿코드"]
    for k in need:
        if k not in col:
            sys.exit(f"양식에 '{k}' 칸이 없습니다.")

    brands = {brand_of(r, col) for r in rows}
    sup = build_supplier(brands)

    R = collections.defaultdict(list)  # 보고서
    out_rows = []
    out_brands = []
    KEEP = "오즈 못 찾은 색상/사이즈(재고 유지)"

    def g(r, k):
        v = r[col[k]]
        return "" if v is None else str(v)

    def lines(r, k):
        v = g(r, k)
        return v.split("\n") if v else []

    for r in rows:
        brand = brand_of(r, col)
        if brand not in sup:
            R["기타 브랜드(업로드 제외)"].append(f"{brand}  {g(r, '상품명')}")
            continue
        name = g(r, "상품명")
        code = g(r, "판매자 상품코드").upper()
        if not code:
            m = CODE_RE.search(name)
            if m:
                code = m.group(1).upper()
                r[col["판매자 상품코드"]] = code
                R["판매자 상품코드 채움"].append(f"{code}  {name}")
        dup_names[(brand, code)] += 1
        is_oz = brand in OZ_BRANDS
        if KIDS_NOTICE_TEMPLATE and (KIDS_RE.search(name) or re.search(r"\d+K$", code)):
            r[col["상품정보제공고시 템플릿코드"]] = KIDS_NOTICE_TEMPLATE
        typ = r[col["옵션형태"]]
        skus = sup[brand].get(code)
        colors_all = ", ".join(sorted({str(s["color"]) for s in skus})) if skus else ""

        # 간편 입력 옵션 -> 엑셀로 수정 불가, 업로드에서 제외
        if not typ:
            R["간편 입력 옵션(엑셀 수정 불가, 업로드 제외)"].append(f"{code}  {name}")
            continue

        # 오즈: 공급처에서 못 찾은 상품은 재고 그대로 (배송비 전환만)
        if not skus and is_oz:
            R["오즈 못 찾은 상품(재고 유지)"].append(f"[{brand}] {code}  {name}")
            out_rows.append(r)
            out_brands.append(brand)
            continue

        # 공급처에 없는 상품 -> 품절
        if not skus:
            if typ == "조합형":
                r[col["옵션 재고수량"]] = "\n".join(["0"] * len(lines(r, "옵션번호")))
            r[col["재고수량"]] = 0
            R["공급처에 없는 상품(품절)"].append(f"[{brand}] {code}  {name}")
            out_rows.append(r)
            out_brands.append(brand)
            continue

        if typ == "설정안함":
            r[col["재고수량"]] = final_qty(sum(s["stock"] for s in skus))
            out_rows.append(r)
            out_brands.append(brand)
            continue

        if typ == "단독형":
            vals = lines(r, "옵션값")
            if len(vals) != 1:
                r[col["재고수량"]] = final_qty(sum(s["stock"] for s in skus))
                R["단독형 2단 옵션(상품 재고만 갱신)"].append(f"{code}  {name}")
                out_rows.append(r)
                out_brands.append(brand)
                continue
            opts = [c.strip() for c in vals[0].split(",") if c.strip()]
            optname = (lines(r, "옵션명") or ["컬러"])[0]
            if is_oz:
                if "사이즈" in optname:
                    miss = [o for o in opts if not size_stock(skus, None, o)[1]]
                else:
                    miss = [o for o in opts if find_color(skus, o, allow_single=False)[0] is None]
                if miss:
                    # 단독형은 옵션별 기존 재고를 알 수 없어 상품 전체를 그대로 둠
                    R[KEEP].append(f"[{brand}] {code}  {', '.join(miss)}  (단독형이라 상품 전체 재고 유지)")
                    out_rows.append(r)
                    out_brands.append(brand)
                    continue
            qtys = []
            if "사이즈" in optname:
                for o in opts:
                    q, ok = size_stock(skus, None, o)
                    if not ok:
                        R["공급처에 없는 사이즈(품절)"].append(f"{code}  '{o}'")
                    qtys.append(final_qty(q))
            else:
                for o in opts:
                    c, guess = find_color(skus, o, allow_guess=(brand != BRAND_DK), allow_single=not is_oz)
                    if c is None:
                        R["공급처에 없는 색상(품절)"].append(f"[{brand}] {code}  '{o}'  (공급처 색상: {colors_all})")
                        qtys.append(0)
                        continue
                    if guess:
                        R["추정 매칭 색상(확인 권장)"].append(f"{code}  '{o}' -> {c}")
                    qtys.append(final_qty(sum(s["stock"] for s in skus if s["color"] == c)))
            r[col["옵션형태"]] = "조합형"
            r[col["옵션명"]] = optname
            r[col["옵션번호"]] = "\n".join(["추가"] * len(opts))
            r[col["옵션값"]] = "\n".join(opts)
            r[col["옵션가"]] = "\n".join(["0"] * len(opts))
            r[col["옵션 재고수량"]] = "\n".join(str(q) for q in qtys)
            r[col["옵션 사용여부"]] = "\n".join(["Y"] * len(opts))
            r[col["재고수량"]] = sum(qtys)
            R["단독형 -> 조합형 전환"].append(f"{code}  " + ", ".join(f"{o}:{q}" for o, q in zip(opts, qtys)))
            out_rows.append(r)
            out_brands.append(brand)
            continue

        # ---------------- 조합형
        optnames = lines(r, "옵션명")
        nums = lines(r, "옵션번호")
        vals = lines(r, "옵션값")
        prices = lines(r, "옵션가")
        uses = lines(r, "옵션 사용여부")
        if len(prices) != len(vals):
            prices = ["0"] * len(vals)
        if len(uses) != len(vals):
            uses = ["Y"] * len(vals)
        size_idx = next((i for i, n in enumerate(optnames) if "사이즈" in n), None)
        color_idx = next((i for i, n in enumerate(optnames) if "사이즈" not in n), None)
        parsed = []
        for v in vals:
            parts = [p.strip() for p in v.split(",")]
            c = parts[color_idx] if color_idx is not None and color_idx < len(parts) else None
            s = parts[size_idx] if size_idx is not None and size_idx < len(parts) else None
            parsed.append((c, s))

        color_cache = {}

        def resolve(c):
            if c is None:
                return None, True
            if c not in color_cache:
                en, guess = find_color(skus, c, allow_guess=(brand != BRAND_DK), allow_single=not is_oz)
                if en is None and is_oz:
                    R[KEEP].append(f"[{brand}] {code}  색상 '{c}'  (공급처 색상: {colors_all})")
                elif en is None:
                    R["공급처에 없는 색상(품절)"].append(f"[{brand}] {code}  '{c}'  (공급처 색상: {colors_all})")
                elif guess:
                    R["추정 매칭 색상(확인 권장)"].append(f"{code}  '{c}' -> {en}")
                color_cache[c] = en
            en = color_cache[c]
            return en, en is not None

        # ---- 사이즈 옵션을 공급처 사이즈표로 다시 맞출지 결정
        rebuild_label = None
        if size_idx is not None:
            if brand == BRAND_DK and any(re.search(r"\(\d", s["size"]) for s in skus):
                rebuild_label = lambda ds: eu_label(ds)
            elif brand == BRAND_EU:
                sup_sizes = {s["size"] for s in skus}
                store_labels = [sz for (_, sz), u in zip(parsed, uses) if sz is not None and u != "N"]
                # 스마트스토어 사이즈 -> 공급처 사이즈 연결
                label_for = {}
                for lab in store_labels:
                    ks = store_size_keys(lab)
                    for ss in sup_sizes:
                        if ks & api_size_keys(ss) and ss not in label_for:
                            label_for[ss] = lab
                unmatched_store = [lab for lab in store_labels
                                   if not any(store_size_keys(lab) & api_size_keys(ss) for ss in sup_sizes)]
                missing_sup = [ss for ss in sup_sizes if ss not in label_for]
                # 스마트스토어 사이즈가 전부 공급처에 있으면 표기 유지 (공급처에만 있는 사이즈는 추가하지 않음)
                if unmatched_store:
                    rebuild_label = lambda ds, lf=label_for: lf.get(ds) or everugg_label(ds)

        if rebuild_label is not None:
            sup_sizes = sorted({s["size"] for s in skus}, key=size_sort_key)
            colors_ko = []
            for c, _ in parsed:
                if c not in colors_ko:
                    colors_ko.append(c)
            new_opts, labels_seen = [], set()
            for c in colors_ko:
                en, ok = resolve(c)
                for ds in sup_sizes:
                    lab = rebuild_label(ds)
                    if (c, lab) in labels_seen:
                        continue
                    labels_seen.add((c, lab))
                    q = sum(s["stock"] for s in skus if s["size"] == ds and (en is None or s["color"] == en)) if ok else 0
                    parts = [None] * len(optnames)
                    if color_idx is not None:
                        parts[color_idx] = c
                    parts[size_idx] = lab
                    new_opts.append((",".join(parts), final_qty(q)))
            old = list(zip(nums, vals, prices))
            used, assign = set(), [None] * len(new_opts)
            for i, (v, _) in enumerate(new_opts):
                for j, (_, ov, _) in enumerate(old):
                    if j not in used and ov == v:
                        assign[i] = j
                        used.add(j)
                        break
            free = [j for j in range(len(old)) if j not in used]
            for i in range(len(new_opts)):
                if assign[i] is None and free:
                    assign[i] = free.pop(0)
            o_num, o_val, o_price, o_qty, o_use = [], [], [], [], []
            for i, (v, q) in enumerate(new_opts):
                j = assign[i]
                o_num.append(old[j][0] if j is not None else "추가")
                o_val.append(v)
                o_price.append(old[j][2] if j is not None else "0")
                o_qty.append(str(q))
                o_use.append("Y")
            for j in free:
                o_num.append(old[j][0])
                o_val.append(old[j][1])
                o_price.append(old[j][2])
                o_qty.append("0")
                o_use.append("N")
            r[col["옵션번호"]] = "\n".join(o_num)
            r[col["옵션값"]] = "\n".join(o_val)
            r[col["옵션가"]] = "\n".join(o_price)
            r[col["옵션 재고수량"]] = "\n".join(o_qty)
            r[col["옵션 사용여부"]] = "\n".join(o_use)
            r[col["재고수량"]] = sum(int(x) for x in o_qty)
            old_sizes = []
            for _, sz in parsed:
                if sz not in old_sizes:
                    old_sizes.append(sz)
            new_sizes = []
            for ds in sup_sizes:
                lab = rebuild_label(ds)
                if lab not in new_sizes:
                    new_sizes.append(lab)
            R["사이즈 재구성"].append(f"[{brand}] {code}  {', '.join(map(str, old_sizes))}  ->  {', '.join(new_sizes)}"
                                    + (f"  (옛 옵션 {len(free)}개 숨김)" if free else ""))
            out_rows.append(r)
            out_brands.append(brand)
            continue

        # 일반: 옵션 표기 유지, 재고만 갱신
        old_q = lines(r, "옵션 재고수량")
        if len(old_q) != len(vals):
            old_q = ["0"] * len(vals)
        qtys = []
        for (c, s), v, use, oq in zip(parsed, vals, uses, old_q):
            en, ok = resolve(c)
            if not ok:
                qtys.append(str(to_int(oq)) if is_oz else "0")
                continue
            if s is None:
                q = sum(x["stock"] for x in skus if en is None or x["color"] == en)
            else:
                q, found = size_stock(skus, en, s)
                if not found and is_oz:
                    if use != "N":
                        sizes = ", ".join(sorted({str(x.get("raw_size", x["size"])) for x in skus
                                                  if en is None or x["color"] == en}))
                        R[KEEP].append(f"[{brand}] {code}  '{v}'  (공급처 사이즈: {sizes})")
                    qtys.append(str(to_int(oq)))
                    continue
                if not found and use != "N":
                    R["공급처에 없는 사이즈(품절)"].append(f"[{brand}] {code}  '{v}'")
            qtys.append(str(final_qty(q)))
        r[col["옵션 재고수량"]] = "\n".join(qtys)
        r[col["재고수량"]] = sum(int(x) for x in qtys)
        out_rows.append(r)
        out_brands.append(brand)

    if PAID_SHIPPING:
        for r in out_rows:
            apply_paid_shipping(r, col, R)

    for (brand, code), n in dup_names.items():
        if n > 1:
            R["중복 등록 상품(확인 권장)"].append(f"[{brand}] {code}  {n}개")

    # 업로드 파일 저장: 첫 번째 양식의 1~2행(제목)을 그대로 쓰고 데이터만 채움
    # 스마트스토어 수정양식 업로드는 한 번에 500개까지라 450개씩 나눠서 저장
    for old in glob.glob("smartstore_stock_upload*.xlsx"):
        os.remove(old)
    out_files = []
    chunks = [out_rows[i:i + UPLOAD_LIMIT] for i in range(0, len(out_rows), UPLOAD_LIMIT)] or [[]]
    src = openpyxl.load_workbook(files[0]).active
    head_rows = [list(r) for r in src.iter_rows(min_row=1, max_row=2, values_only=True)]
    head_merges = [str(m) for m in src.merged_cells.ranges if m.max_row <= 2]
    for n, chunk in enumerate(chunks, start=1):
        # 양식 파일을 그대로 쓰면 빈 서식 행이 남아서 "500개 초과" 오류가 나므로 새 파일에 값만 씀
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = src.title
        for r in head_rows + chunk:
            ws.append(r)
        for m in head_merges:
            ws.merge_cells(m)
        name = OUT_FILE if len(chunks) == 1 else OUT_FILE.replace(".xlsx", f"_{n}.xlsx")
        wb.save(name)
        out_files.append(f"{name} ({len(chunk)}개)")

    order = ["오즈 못 찾은 상품(재고 유지)", KEEP, "유료배송 전환", "배송비 유형 확인 필요(변경 안 함)", "사이즈 재구성", "공급처에 없는 상품(품절)", "공급처에 없는 색상(품절)", "공급처에 없는 사이즈(품절)",
             "추정 매칭 색상(확인 권장)", "중복 등록 상품(확인 권장)",
             "간편 입력 옵션(엑셀 수정 불가, 업로드 제외)", "단독형 2단 옵션(상품 재고만 갱신)",
             "단독형 -> 조합형 전환", "판매자 상품코드 채움", "기타 브랜드(업로드 제외)"]
    cnt = collections.Counter(out_brands)
    summary = [f"⚠️ {n}" for n in SUPPLIER_NOTES] + [f"업로드 파일: {len(out_files)}개 (" + ", ".join(out_files) + ")",
               f"업로드 파일 상품 수: {len(out_rows)}  ({', '.join(f'{k} {v}' for k, v in cnt.items())})",
               f"안전재고: {SAFETY_STOCK}"]
    if BRAND_OL in sup or BRAND_OW in sup:
        summary.append(f"오즈 사이즈 변환: AU 숫자 + {OZ_EU_OFFSET} = EU (예: AU5 -> EU{5 + OZ_EU_OFFSET})")
    summary += [f"{k}: {len(R[k])}" for k in order if R[k]]
    body = []
    for k in order:
        if R[k]:
            body.append(f"\n■ {k} ({len(R[k])})")
            body += ["  " + x for x in R[k]]
    with open(REPORT_FILE, "w", encoding="utf-8") as f:
        f.write("\n".join(summary) + "\n" + "\n".join(body) + "\n")
    print("\n".join(summary))
    print("\n완료:", ", ".join(out_files), REPORT_FILE)


if __name__ == "__main__":
    main()
