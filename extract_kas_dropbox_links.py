"""
K-AS UGG 가격표 엑셀에서 스타일코드 + White BG 드랍박스 링크를 뽑아서
apparel_dropbox_links.csv 형식(code,dropbox_folder_url)으로 저장합니다.
fetch_dropbox_style_images.py에 바로 이어서 쓸 수 있습니다.

사용법:
    python extract_kas_dropbox_links.py --xlsx "K-AS_UGG_Price_List_18_08_2026.xlsx" --sheet "2026" --output "apparel_dropbox_links.csv"
"""

import argparse
import csv

import openpyxl

STYLE_HEADER_CANDIDATES = ["STYLE NO.", "STYLE NO"]
WHITEBG_HEADER_CANDIDATES = ["White BG"]


def find_header_row(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
        if row and any(isinstance(v, str) and v.strip().upper().startswith("STYLE NO") for v in row if v):
            return i
    raise RuntimeError("헤더 행(STYLE NO.)을 찾지 못했습니다.")


def build_col_index(headers):
    col_index = {}
    for idx, h in enumerate(headers):
        if isinstance(h, str):
            col_index[h.strip()] = idx
    return col_index


def get_first(row, col_index, candidates):
    for name in candidates:
        idx = col_index.get(name)
        if idx is not None and idx < len(row) and row[idx] not in (None, ""):
            return row[idx]
    return None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--output", default="apparel_dropbox_links.csv")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[args.sheet]

    header_row_idx = find_header_row(ws)
    headers = [c.value for c in ws[header_row_idx]]
    col_index = build_col_index(headers)

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        code = get_first(row, col_index, STYLE_HEADER_CANDIDATES)
        url = get_first(row, col_index, WHITEBG_HEADER_CANDIDATES)
        if code and isinstance(code, str) and code.strip() and url:
            rows.append((code.strip(), str(url).strip()))

    with open(args.output, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["code", "dropbox_folder_url"])
        writer.writerows(rows)

    print(f"완료: {args.output} 생성 ({len(rows)}개 코드)")


if __name__ == "__main__":
    main()
